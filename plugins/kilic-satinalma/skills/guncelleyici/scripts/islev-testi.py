"""Functional fixture tests; never opens user documents or a pre-existing Excel session."""
from __future__ import annotations
import importlib
import json
from pathlib import Path
import sys
import time


def emit(name, status, detail):
    print("@@T " + json.dumps({"name": name, "status": status, "detail": str(detail)}, ensure_ascii=False), flush=True)


def pdf_fixture(path):
    content = b"BT /F1 20 Tf 72 760 Td (GUNCELLEYICI TEST 12345) Tj ET"
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
               b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
               b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
               b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    data = bytearray(b"%PDF-1.4\n"); offsets = []
    for i, obj in enumerate(objects, 1):
        offsets.append(len(data)); data += b"%d 0 obj\n" % i + obj + b"\nendobj\n"
    xref = len(data); data += b"xref\n0 6\n0000000000 65535 f \n"
    for offset in offsets: data += b"%010d 00000 n \n" % offset
    data += b"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % xref
    Path(path).write_bytes(data)


def excel_test(path):
    import pythoncom
    import win32api
    import win32con
    import win32event
    import win32process
    import win32com.client
    pythoncom.CoInitialize()
    xl = handle = None
    independent = False
    existing = set(win32process.EnumProcesses())
    try:
        xl = win32com.client.DispatchEx("Excel.Application")
        pid = win32process.GetWindowThreadProcessId(xl.Hwnd)[1]
        if not pid or pid in existing:
            raise RuntimeError("Ayrı test Excel süreci doğrulanamadı; mevcut oturuma dokunulmadı.")
        handle = win32api.OpenProcess(win32con.SYNCHRONIZE | win32con.PROCESS_TERMINATE, False, pid)
        independent = True
        xl.Visible = False; xl.DisplayAlerts = False; xl.ScreenUpdating = False
        wb = xl.Workbooks.Open(str(path), 0, True)
        try:
            xl.CalculateFull()
            value = wb.Worksheets(1).Range("A2").Value
            if value != 42: raise RuntimeError(f"Excel hesap sonucu farklı: {value!r}")
        finally:
            wb.Close(False); del wb
        emit("Excel calculate", "PASS", "21*2=42")
        Path(str(path) + ":Zone.Identifier").write_text("[ZoneTransfer]\nZoneId=3\n", encoding="ascii")
        try:
            wb = xl.Workbooks.Open(str(path), 0, True)
            try:
                if wb.Worksheets(1).Range("A2").Value != 42: raise RuntimeError("Internet-zone workbook result mismatch")
            finally: wb.Close(False); del wb
            emit("Excel internet-zone fixture", "PASS", "Kaynak fixture değişmeden açıldı")
        except Exception as exc:
            emit("Excel internet-zone fixture", "WARNING", exc)
    finally:
        if xl is not None:
            if independent:
                try: xl.Quit()
                except Exception: pass
            del xl
        pythoncom.CoUninitialize()
        if handle is not None:
            try:
                if win32event.WaitForSingleObject(handle, 20000) != win32event.WAIT_OBJECT_0:
                    win32api.TerminateProcess(handle, 1)
                    emit("Excel cleanup", "WARNING", "Yalnız testin açtığı süreç kapatıldı")
                else: emit("Excel cleanup", "PASS", "Test süreci kapandı")
            finally: win32api.CloseHandle(handle)


def main():
    root = Path(sys.argv[1]).resolve(); root.mkdir(parents=True, exist_ok=True)
    pdf, xlsx = root / "test.pdf", root / "test.xlsx"
    pdf_fixture(pdf)
    loaded = set()
    for module in ("openpyxl", "pypdf", "pdfplumber", "pypdfium2", "PIL.Image", "cryptography.hazmat.primitives.hashes", "pythoncom", "win32com.client"):
        try: importlib.import_module(module); loaded.add(module); emit(module, "PASS", "import")
        except Exception as exc: emit(module, "FAIL", exc)
    tests = []
    def xlsx_test():
        import openpyxl
        wb = openpyxl.Workbook(); ws = wb.active; ws["A1"] = 21; ws["A2"] = "=A1*2"; wb.save(xlsx); wb.close()
        back = openpyxl.load_workbook(xlsx)
        try:
            if back.active["A2"].value != "=A1*2": raise RuntimeError("Formula missing")
        finally: back.close()
    def pypdf_test():
        from pypdf import PdfReader
        if "12345" not in (PdfReader(pdf).pages[0].extract_text() or ""): raise RuntimeError("PDF text missing")
    def plumber_test():
        import pdfplumber
        with pdfplumber.open(pdf) as doc:
            if "12345" not in (doc.pages[0].extract_text() or ""): raise RuntimeError("PDF text missing")
    def render_test():
        import pypdfium2
        with pypdfium2.PdfDocument(pdf) as doc:
            page = doc[0]; bitmap = page.render(scale=1)
            try:
                image = bitmap.to_pil(); image.save(root / "page.png")
                if image.width < 100: raise RuntimeError("Render dimensions invalid")
                image.close()
            finally: bitmap.close(); page.close()
    def crypto_test():
        from cryptography.hazmat.primitives import hashes
        import hashlib
        h = hashes.Hash(hashes.SHA256()); h.update(b"fixture")
        if h.finalize() != hashlib.sha256(b"fixture").digest(): raise RuntimeError("Hash mismatch")
    tests = [("openpyxl", xlsx_test), ("pypdf", pypdf_test), ("pdfplumber", plumber_test),
             ("pypdfium2", render_test), ("cryptography.hazmat.primitives.hashes", crypto_test)]
    for module, fn in tests:
        if module not in loaded:
            emit(module + " functional", "NOT_RUN", "Import failed"); continue
        try: fn(); emit(module + " functional", "PASS", "Fixture doğrulandı")
        except Exception as exc: emit(module + " functional", "FAIL", exc)
    if "pythoncom" in loaded and "win32com.client" in loaded and xlsx.exists():
        try: excel_test(xlsx)
        except Exception as exc: emit("Excel", "FAIL", exc)
    else: emit("Excel", "NOT_RUN", "Windows/pywin32/fixture hazır değil")
    return 0


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"): sys.stdout.reconfigure(encoding="utf-8")
    raise SystemExit(main())